"""PPTX extraction: shapes (including grouped shapes), tables, speaker notes, document
properties and picture OCR. Positions are in EMU, pages are slide numbers."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from ..config import Settings
from ..models import Document
from .common import IdGen
from .docx import count_structure, ocr_part, unit_to_span
from .ooxml import iter_pictures, package_units, pptx_units


def extract_pptx(path: str | Path, settings: Settings) -> Document:
    path = Path(path)
    doc = Document(file=path.name, path=str(path), file_type="pptx")
    ids = IdGen(path.stem[:12])
    prs = Presentation(str(path))
    doc.pages = len(prs.slides)
    for i in range(1, doc.pages + 1):
        doc.page_sizes[i] = (prs.slide_width, prs.slide_height)

    units = list(pptx_units(prs)) + list(package_units(prs.part.package))
    doc.spans.extend(unit_to_span(u, ids(), doc.file) for u in units)
    counts = count_structure(units)
    counts["slides"] = doc.pages

    if settings.ocr_embedded_images:
        seen: set[str] = set()
        for sno, slide in enumerate(prs.slides, start=1):
            for pic, anchor in iter_pictures(slide.shapes, f"slide[{sno}]"):
                try:
                    part = pic.image.part if hasattr(pic.image, "part") else slide.part.related_part(pic._element.blip_rId)
                except Exception:
                    continue
                key = f"{part.partname}@{sno}"
                if key in seen:
                    continue
                seen.add(key)
                bbox = (pic.left, pic.top, pic.left + pic.width, pic.top + pic.height)
                doc.spans.extend(ocr_part(part, doc, ids, settings, page=sno,
                                          location=f"slide {sno}, picture '{pic.name}'", bbox=bbox))
    counts["images"] = len(doc.images)
    doc.structure = counts
    return doc
