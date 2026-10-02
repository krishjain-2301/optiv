"""Planted PII in the places a text walker never looks: hyperlink targets, field codes, tracked
deletions, document properties, alt text, image metadata, chart caches, embedded workbooks and
thumbnails. Every output is searched for every planted value, in every part, recursively."""
import html
import io
import zipfile
from pathlib import Path
from urllib.parse import unquote

import pytest
from lxml import etree
from PIL import Image
from PIL.PngImagePlugin import PngInfo

from pii_shield.extract.ooxml import W

PLANTED_DOCX = {
    "mailto": "kofi.mensah@example.com",
    "field": "rahul.verma@example.com",
    "deleted": "Siobhán Walsh",
    "description": "Priya Raman",
    "alt": "Larry Page",
    "exif": "Hans Müller",
}
PLANTED_PPTX = {
    "chart1": "Hans Müller",
    "chart2": "Łukasz Nowak",
    "alt": "Larry Page",
    "description": "Priya Raman",
    "caps": "GRACE WEST",
}


def _png(author: str | None = None, size=(200, 200)) -> bytes:
    img = Image.new("RGB", size, "white")
    buf = io.BytesIO()
    info = None
    if author:
        info = PngInfo()
        info.add_text("Author", author)
    img.save(buf, format="PNG", pnginfo=info)
    return buf.getvalue()


def make_docx(path: Path) -> None:
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    d = docx.Document()
    d.add_heading("Vendor onboarding notes", 1)
    p = d.add_paragraph("Questions about this procedure: ")
    rid = d.part.relate_to(f"mailto:{PLANTED_DOCX['mailto']}", RT.HYPERLINK, is_external=True)
    p._p.append(etree.fromstring(
        f'<w:hyperlink xmlns:w="{W}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{rid}"><w:r><w:t>email the owner</w:t></w:r></w:hyperlink>'))
    p2 = d.add_paragraph("Escalations go to ")
    p2._p.append(etree.fromstring(
        f'<w:fldSimple xmlns:w="{W}" w:instr=\' HYPERLINK "mailto:{PLANTED_DOCX["field"]}" \'>'
        f'<w:r><w:t>the escalation desk</w:t></w:r></w:fldSimple>'))
    p3 = d.add_paragraph("Final approval is pending.")
    p3._p.append(etree.fromstring(
        f'<w:del xmlns:w="{W}" w:id="1" w:author="Reviewer" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:delText>Approved by {PLANTED_DOCX["deleted"]}</w:delText></w:r></w:del>'))
    d.add_picture(io.BytesIO(_png(PLANTED_DOCX["exif"])))
    d.inline_shapes[0]._inline.docPr.set("descr", f"Photo of {PLANTED_DOCX['alt']}")
    d.core_properties.comments = f"Drafted for {PLANTED_DOCX['description']} by the risk office"
    buf = io.BytesIO()
    d.save(buf)
    # A thumbnail, as Word writes one when "save preview picture" is on.
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "_rels/.rels":
                data = data.replace(b"</Relationships>",
                                    b'<Relationship Id="rIdThumb" Type="http://schemas.openxmlformats.org/package/2006/'
                                    b'relationships/metadata/thumbnail" Target="docProps/thumbnail.png"/></Relationships>')
            out.writestr(info, data)
        out.writestr("docProps/thumbnail.png", _png(PLANTED_DOCX["exif"], (64, 64)))


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Open actions by owner"
    data = CategoryChartData()
    data.categories = [PLANTED_PPTX["chart1"], PLANTED_PPTX["chart2"]]
    data.add_series("Open actions", (4, 7))
    s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1), Inches(1.5), Inches(6), Inches(4), data)
    pic = s.shapes.add_picture(io.BytesIO(_png()), Inches(7), Inches(1.5), Inches(2), Inches(2))
    pic._element.nvPicPr.cNvPr.set("descr", f"Photo of {PLANTED_PPTX['alt']}")
    tb = s.shapes.add_textbox(Inches(1), Inches(6), Inches(6), Inches(1))
    tb.text_frame.text = f"Chart prepared by {PLANTED_PPTX['caps']} for the risk committee."
    prs.core_properties.comments = f"Drafted for {PLANTED_PPTX['description']} by the risk office"
    prs.save(path)


def members(data: bytes, prefix=""):
    """Every member of a package, nested packages included."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for n in zf.namelist():
            body = zf.read(n)
            if body[:4] == b"PK\x03\x04":
                yield from members(body, f"{prefix}{n}!")
            else:
                yield prefix + n, body


def views(body: bytes) -> str:
    t = body.decode("utf-8", errors="ignore") + "\n" + body.decode("latin-1")
    stripped = etree.XML  # noqa: F841  (only for readability below)
    import re

    no_tags = re.sub(r"<[^>]+>", "", t)
    return "\n".join({t, no_tags, html.unescape(t), html.unescape(no_tags), unquote(html.unescape(t))}).lower()


def surviving(path: Path, values) -> list[tuple[str, str]]:
    found = []
    for name, body in members(path.read_bytes()):
        v = views(body)
        found += [(name, x) for x in values if x.lower() in v]
    return found


@pytest.fixture(scope="module")
def planted_run(tmp_path_factory, settings):
    from pii_shield import run

    src = tmp_path_factory.mktemp("planted")
    make_docx(src / "planted.docx")
    make_pptx(src / "planted.pptx")
    out = tmp_path_factory.mktemp("planted_out")
    res = run([src / "planted.docx", src / "planted.pptx"], settings, out)
    res.out_dir, res.src = out, src
    return res


def test_planted_sources_really_contain_the_values(planted_run):
    # Guards the test itself: every planted value is present in the input somewhere.
    assert {v for _, v in surviving(planted_run.src / "planted.docx", PLANTED_DOCX.values())} == set(PLANTED_DOCX.values())
    assert {v for _, v in surviving(planted_run.src / "planted.pptx", PLANTED_PPTX.values())} == set(PLANTED_PPTX.values())


def test_masked_files_written(planted_run):
    assert not planted_run.errors, planted_run.errors
    assert (planted_run.out_dir / "planted.masked.docx").exists()
    assert (planted_run.out_dir / "planted.masked.pptx").exists()


@pytest.mark.parametrize("name,planted", [("planted.masked.docx", PLANTED_DOCX), ("planted.masked.pptx", PLANTED_PPTX)])
def test_no_planted_value_in_any_part(planted_run, name, planted):
    left = surviving(planted_run.out_dir / name, planted.values())
    assert not left, left


@pytest.mark.parametrize("name", ["planted.masked.docx", "planted.masked.pptx"])
def test_thumbnail_and_embeddings_removed(planted_run, name):
    names = [n for n, _ in members((planted_run.out_dir / name).read_bytes())]
    assert not [n for n in names if "thumbnail" in n or "/embeddings/" in n or "!" in n], names


def test_llm_text_has_no_planted_value(planted_run):
    for f, text in planted_run.redacted.items():
        planted = PLANTED_DOCX if f.endswith("docx") else PLANTED_PPTX
        low = text.lower()
        assert not [v for v in planted.values() if v.lower() in low], (f, text)


def test_chart_and_description_reach_the_llm_text_redacted(planted_run):
    md = planted_run.redacted["planted.pptx"]
    assert "[chart]" in md and "[PERSON_" in md
    assert "Drafted for [PERSON_" in md


def test_gate_refuses_a_package_with_a_surviving_value(tmp_path):
    from pii_shield.redact.leakcheck import LeakError, build_needles, check_package, scrub_package
    from pii_shield.redact.tokens import TokenVault

    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as z:
        z.writestr("word/styles.xml", f'<w:styles xmlns:w="{W}"><w:style w:styleId="x"><w:name w:val="Ines Duarte"/></w:style></w:styles>')
    v = TokenVault()
    v.values["[PERSON_001]"].add("Ines Duarte")
    n = build_needles(v)
    scrubbed, _ = scrub_package(data.getvalue(), n)  # w:val is not a place the scrub may rewrite
    with pytest.raises(LeakError):
        check_package(scrubbed, n, "test")


def test_gate_finds_values_split_across_runs_and_escaped():
    from pii_shield.redact.leakcheck import build_needles, scan_bytes
    from pii_shield.redact.tokens import TokenVault

    v = TokenVault()
    v.values["[PERSON_001]"].add("Zoë O'Brien")
    v.values["[EMAIL_001]"].add("zoe.obrien@example.com")
    v.values["[PHONE_001]"].add("+1 (212) 555-0193")
    n = build_needles(v)
    assert scan_bytes("a.xml", f"<w:t>Zoë O'Bri</w:t></w:r><w:r><w:t>en</w:t>".encode(), n)
    assert scan_bytes("a.xml", "<w:t>Zoë O&apos;Brien</w:t>".encode(), n)
    assert scan_bytes("a.rels", b'Target="mailto:zoe.obrien%40example.com"', n)
    assert scan_bytes("a.rels", b'Target="tel:+12125550193"', n)
    assert not scan_bytes("a.xml", b"<w:t>nothing to see</w:t>", n)
