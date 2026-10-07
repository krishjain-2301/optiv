"""What the leak gate could not see, and what the pipeline now does about it: pixels are read
again after masking, pictures nobody read are blanked, a reviewer's decisions rewrite every
output, the audit log is a hash chain, and tokens can be keyed and put back."""
import io
import json
import zipfile

import numpy as np
import pymupdf as fitz
import pytest
from PIL import Image, ImageDraw

from optiv_pii_shield import Settings, audit, run
from optiv_pii_shield.models import Document, Finding, Span
from optiv_pii_shield.redact.leakcheck import LeakError, build_needles, check_pdf
from optiv_pii_shield.redact.tokens import TokenVault, rehydrate

from make_samples import font

VALUES = {"[PERSON_001]": "Rafael Mendoza-Kowalski", "[SSN_001]": "900-12-3456"}


def needles_for(values=VALUES):
    v = TokenVault()
    for token, value in values.items():
        v.values[token].add(value)
    return build_needles(v)


def text_picture(lines, size=(1600, 500)) -> Image.Image:
    img = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        d.text((60, 60 + i * 90), line, fill="black", font=font(44))
    return img


def scanned_pdf(path, lines):
    buf = io.BytesIO()
    text_picture(lines).save(buf, format="PNG")
    pdf = fitz.open()
    page = pdf.new_page(width=576, height=180)
    page.insert_image(page.rect, stream=buf.getvalue())
    pdf.save(path)
    return Document(file=path.name, path=str(path), file_type="pdf", pages=1, ocr_pages=[1])


# ------------------------------------------------------------------- verification pass
def test_text_gate_alone_cannot_see_a_scan_but_verification_can(tmp_path):
    from optiv_pii_shield.redact import verify

    path = tmp_path / "scan.pdf"
    doc = scanned_pdf(path, ["Full name: Rafael Mendoza-Kowalski", "SSN: 900-12-3456  Status: active"])
    n = needles_for()
    check_pdf(path, n, "scan")  # passes: the values are pixels, and this check reads text
    summary = verify.verify_pdf(doc, path, Settings(), n)
    assert summary["covered"] >= 2  # read again as a picture, found, covered
    again = verify.verify_pdf(doc, path, Settings(), n)
    assert again["covered"] == 0
    with fitz.open(path) as pdf:  # the rest of the page is still there
        pix = pdf[0].get_pixmap(dpi=150)
    assert 0.02 < (np.frombuffer(pix.samples, np.uint8) < 80).mean() < 0.6


def test_verification_withholds_what_it_cannot_cover(tmp_path, monkeypatch):
    from optiv_pii_shield.redact import verify

    path = tmp_path / "scan.pdf"
    doc = scanned_pdf(path, ["SSN: 900-12-3456"])
    monkeypatch.setattr(verify, "MAX_ROUNDS", 0)  # no attempt to cover: the value is still readable
    with pytest.raises(LeakError):
        verify.verify_pdf(doc, path, Settings(), needles_for())


def test_verification_covers_a_value_left_in_a_picture(tmp_path):
    from optiv_pii_shield.redact import verify

    path = tmp_path / "shot.png"
    text_picture(["Owner: Rafael Mendoza-Kowalski", "Ticket 4471 closed"]).save(path)
    doc = Document(file="shot.png", path=str(path), file_type="image", pages=1)
    assert verify.verify_image(doc, path, Settings(), needles_for())["covered"] >= 1
    assert verify.verify_image(doc, path, Settings(), needles_for())["covered"] == 0


def test_verification_blanks_a_package_picture_that_still_shows_a_value():
    from optiv_pii_shield.redact import verify

    shot, clean = io.BytesIO(), io.BytesIO()
    text_picture(["Reviewer: Rafael Mendoza-Kowalski"]).save(shot, format="PNG")
    text_picture(["Quarterly vendor review"]).save(clean, format="PNG")
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as z:
        z.writestr("word/document.xml", "<w/>")
        z.writestr("word/media/image1.png", shot.getvalue())
        z.writestr("word/media/image2.png", clean.getvalue())
    doc = Document(file="a.docx", path="a.docx", file_type="docx")
    new, summary = verify.verify_package(doc, data.getvalue(), Settings(), needles_for())
    assert summary == {"method": "re-OCR", "pictures": 2, "covered": 1}
    with zipfile.ZipFile(io.BytesIO(new)) as z:
        assert z.read("word/media/image1.png") != shot.getvalue()
        assert z.read("word/media/image2.png") == clean.getvalue()


def test_run_reports_verification(run_result):
    for doc in run_result.docs.values():
        assert doc.gate["verified"]["method"] == "re-OCR", doc.file
    assert run_result.docs["risk_policy_scanned.pdf"].gate["verified"]["pages"] == 3


# ------------------------------------------------- the gate's catches become findings
def test_a_value_found_once_is_redacted_everywhere_it_appears():
    from optiv_pii_shield.pipeline import GATE_LAYER, locate_missed

    spans = [Span(id="a", file="f", text="Ref QX-99-ZETA-7 opened", kind="paragraph", page=1),
             Span(id="b", file="f", text="see qx-99-zeta-7 again", kind="paragraph", page=1)]
    doc = Document(file="f", path="f", file_type="text", spans=spans)
    found = Finding("a", "f", 4, 16, "QX-99-ZETA-7", "EMPLOYEE_ID", 0.9, "r", "L1 rules")
    findings = {"f": [found]}
    vault = TokenVault()
    vault.assign_all({"f": doc}, findings)
    assert locate_missed({"f": doc}, findings, build_needles(vault, findings), vault) == 1
    extra = [f for f in findings["f"] if f.layer == GATE_LAYER]
    assert [(f.span_id, f.text, f.token, f.entity_type) for f in extra] == [("b", "qx-99-zeta-7", found.token, "EMPLOYEE_ID")]


# ------------------------------------------------------------- pictures nobody could read
def photo_png(size=(240, 240)) -> bytes:
    rng = np.random.default_rng(7)
    arr = (rng.random((size[1] // 8, size[0] // 8, 3)) * 255).astype(np.uint8)
    buf = io.BytesIO()
    Image.fromarray(arr).resize(size, Image.BICUBIC).save(buf, format="PNG")
    return buf.getvalue()


def docx_with_photo(path):
    import docx

    d = docx.Document()
    d.add_paragraph("Badge photo of the reviewer:")
    d.add_picture(io.BytesIO(photo_png()))
    d.save(path)


def media(path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist() if "/media/" in n}


@pytest.mark.parametrize("change,blanked", [({}, True), ({"blank_textless_images": False}, False),
                                            ({"ocr_embedded_images": False}, True)])
def test_pictures_without_readable_text_are_blanked(tmp_path, change, blanked):
    src = tmp_path / "badge.docx"
    docx_with_photo(src)
    res = run([src], Settings(**change), tmp_path / "out")
    assert not res.errors, res.errors
    before, after = media(src), media(tmp_path / "out" / "badge.masked.docx")
    assert len(after) == 1
    same = list(before.values())[0] == list(after.values())[0]
    assert same != blanked
    if blanked:
        img = np.array(Image.open(io.BytesIO(list(after.values())[0])).convert("RGB"))
        assert np.median(img) == 60  # the "IMAGE WITHHELD" grey


def test_qr_codes_are_found_and_models_are_pinned(tmp_path, monkeypatch):
    import cv2

    from optiv_pii_shield import modelstore
    from optiv_pii_shield.errors import ModelMissing
    from optiv_pii_shield.extract import visual

    qr = cv2.QRCodeEncoder.create().encode("BEGIN:VCARD FN:Priya Raman TEL:+919820055512 END:VCARD")
    qr = cv2.resize(qr, (300, 300), interpolation=cv2.INTER_NEAREST)
    page = np.full((800, 1000), 255, np.uint8)
    page[200:500, 300:600] = qr
    found = visual.detect(page, Settings())
    assert [k for k, _ in found] == ["qr"]
    x0, y0, x1, y1 = found[0][1]
    assert x0 < 330 and y0 < 230 and x1 > 570 and y1 > 470
    assert visual.detect(np.full((400, 400, 3), 255, np.uint8), Settings()) == []  # loads the face model too
    # a model file that is not the pinned one is refused
    monkeypatch.setattr(modelstore, "MODELS_DIR", tmp_path)
    (tmp_path / modelstore.FACE.name).write_bytes(b"not the model")
    with pytest.raises(ModelMissing):
        modelstore.locate(modelstore.FACE)


def test_pdf_bookmarks_are_redacted(tmp_path):
    src = tmp_path / "book.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 100), "Appendix C. Personnel record. Full name: Rafael Mendoza-Kowalski, SSN: 900-12-3456.", fontsize=11)
    pdf.set_toc([[1, "Record of Rafael Mendoza-Kowalski", 1]])
    pdf.save(src)
    res = run([src], Settings(), tmp_path / "out")
    assert not res.errors, res.errors
    with fitz.open(tmp_path / "out" / "book.masked.pdf") as out:
        title = out.get_toc()[0][1]
        assert "Rafael" not in title and "Mendoza" not in title and "[PERSON_" in title
        assert "Rafael" not in out[0].get_text()


# ------------------------------------------------------------------- formats and inputs
@pytest.fixture(scope="module")
def text_run(tmp_path_factory):
    src = tmp_path_factory.mktemp("plain")
    (src / "notes.txt").write_text(
        "Full name: Priya Raman\n\nThe Cloud migration was approved. Escalate to priya.raman@example.com.\n\n"
        "Afterwards zorblat quux signed the form and paid from account 123456789012 at HDFC0001234.\n", encoding="utf-8")
    (src / "people.csv").write_text(
        "Name,E-mail,Blood group,UPI\nHans Müller,hans.mueller@example.com,O+,hans.m@okhdfcbank\n", encoding="utf-8")
    (src / "mail.eml").write_bytes(
        b"From: Grace Wilson <grace.wilson@example.com>\r\nTo: Tomasz Zielinski <t.zielinski@example.com>\r\n"
        b"Subject: Access for Priya Raman\r\nDate: Mon, 5 Oct 2026 09:00:00 +0000\r\nMIME-Version: 1.0\r\n"
        b"Content-Type: multipart/mixed; boundary=XX\r\n\r\n--XX\r\nContent-Type: text/plain; charset=utf-8\r\n\r\n"
        b"Please call Priya Raman on +44 7700 900418 about the renewal.\r\n--XX\r\n"
        b"Content-Type: application/pdf\r\nContent-Disposition: attachment; filename=cv.pdf\r\n\r\n%PDF-fake\r\n--XX--\r\n")
    (src / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)
    other = src / "second"
    other.mkdir()
    (other / "notes.txt").write_text("Contact Kofi Mensah at kofi.mensah@example.com.\n", encoding="utf-8")
    files = [src / "notes.txt", src / "people.csv", src / "mail.eml", src / "old.doc", other / "notes.txt"]
    res = run(files, Settings(operator="tester"), src / "out")
    res.src = src
    return res


def test_text_csv_and_eml_are_read_and_masked(text_run):
    out = text_run.out_dir
    assert set(text_run.docs) == {"notes.txt", "people.csv", "mail.eml", "notes~2.txt"}  # same name twice: both kept
    assert "legacy Office" in text_run.errors["old.doc"]  # refused, reported, nothing passed on
    planted = ["Priya Raman", "priya.raman@example.com", "Hans Müller", "hans.mueller@example.com", "hans.m@okhdfcbank",
               "Grace Wilson", "grace.wilson@example.com", "Tomasz Zielinski", "+44 7700 900418", "123456789012",
               "HDFC0001234", "Kofi Mensah"]
    for name in ("notes.masked.txt", "people.masked.csv", "mail.masked.eml", "notes~2.masked.txt"):
        text = (out / name).read_text(encoding="utf-8")
        assert not [v for v in planted if v.lower() in text.lower()], (name, text)
    csv_text = (out / "people.masked.csv").read_text(encoding="utf-8")
    assert csv_text.splitlines()[0] == "Name,E-mail,Blood group,UPI" and "[HEALTH" in csv_text and "[UPI" in csv_text
    mail = (out / "mail.masked.eml").read_text(encoding="utf-8")
    assert "From: [PERSON_" in mail and "1 attachment(s) removed" in mail and "%PDF" not in mail
    assert any("attachment" in w for w in text_run.docs["mail.eml"].warnings)


# ------------------------------------------------------------------------------- review
def test_review_decisions_rewrite_every_output(text_run):
    from optiv_pii_shield import review
    from optiv_pii_shield.pipeline import apply_review

    out = text_run.out_dir
    notes = (out / "notes.txt.redacted.md").read_text(encoding="utf-8")
    assert "zorblat quux" in notes  # a name no layer has evidence for: the kind of miss a reviewer adds
    live = [f for f in text_run.findings["notes.txt"] if f.decision != "drop"]
    email = next(f for f in live if f.entity_type == "EMAIL_ADDRESS")
    counts = apply_review(text_run, [review.Decision("EMAIL_ADDRESS", email.text, "reject")],
                          [review.Addition("Zorblat Quux", "PERSON")], operator="reviewer-1")
    assert counts == {"approved": 0, "rejected": 1, "added": 1}
    notes = (out / "notes.txt.redacted.md").read_text(encoding="utf-8")
    masked = (out / "notes.masked.txt").read_text(encoding="utf-8")
    for text in (notes, masked):
        assert "zorblat" not in text.lower() and "quux" not in text.lower() and email.text in text
    assert "Priya Raman" not in masked  # earlier findings still hold
    # the audit chain goes on: the review and the changed findings are appended, nothing rewritten
    records = [json.loads(l) for l in (out / "audit_log.jsonl").read_text(encoding="utf-8").splitlines()]
    events = [r["event"] for r in records]
    assert events[0] == "run_started" and events.count("review_applied") == 1 and events.count("outputs_written") == 2
    applied = next(r for r in records if r["event"] == "review_applied")
    assert applied["operator"] == "reviewer-1" and applied["added"] == 1
    assert "zorblat" not in (out / "audit_log.jsonl").read_text(encoding="utf-8").lower()
    assert audit.verify_run(out)["ok"]
    # reviewed findings are marked as such in a gold draft
    from optiv_pii_shield.evaluate import gold_template

    draft = gold_template(text_run.findings, out / "draft.SENSITIVE.csv").read_text(encoding="utf-8")
    assert "zorblat quux,PERSON,narrative,reviewer: added" in draft


# ------------------------------------------------------------------ audit and manifest
def test_manifest_and_audit_chain_detect_changes(tmp_path):
    src = tmp_path / "a.txt"
    src.write_text("Full name: Priya Raman\n\nSSN: 900-12-3456\n", encoding="utf-8")
    out = tmp_path / "out"
    run([src], Settings(operator="alice"), out)
    manifest = json.loads((out / "run_manifest.json").read_text(encoding="utf-8"))
    assert manifest["operator"] == "alice" and manifest["inputs"][0]["sha256"] == audit.modelstore.sha256_file(src)
    assert manifest["settings"]["vault_passphrase"] == "not set" and manifest["versions"]["spacy"]
    assert manifest["signature"] is None
    assert "Priya" not in json.dumps(manifest)
    ok = audit.verify_run(out)
    assert ok["ok"] and not ok["signed"] and ok["audit"]["records"] == manifest["audit"]["records"]

    log = out / "audit_log.jsonl"
    original = log.read_text(encoding="utf-8")
    lines = original.splitlines()
    log.write_text("\n".join(lines[:1] + lines[2:]) + "\n", encoding="utf-8")  # a record removed
    assert any("audit log" in p for p in audit.verify_run(out)["problems"])
    log.write_text(original.replace('"decision": "redact"', '"decision": "drop"', 1), encoding="utf-8")  # a record edited
    assert any("changed after" in p for p in audit.verify_run(out)["problems"])
    log.write_text(original, encoding="utf-8")
    red = out / "a.txt.redacted.md"
    red.write_text(red.read_text(encoding="utf-8") + "edited", encoding="utf-8")  # an output edited
    assert any("a.txt.redacted.md" in p for p in audit.verify_run(out)["problems"])


def test_signed_manifest(tmp_path, monkeypatch):
    priv, pub = audit.keygen(tmp_path / "keys")
    monkeypatch.setenv(audit.SIGNING_KEY_ENV, str(priv))
    src = tmp_path / "a.txt"
    src.write_text("Contact kofi.mensah@example.com\n", encoding="utf-8")
    run([src], Settings(), tmp_path / "out")
    key = pub.read_text(encoding="ascii")
    assert audit.verify_run(tmp_path / "out", key) == {**audit.verify_run(tmp_path / "out", key), "ok": True, "signed": True}
    other = audit.keygen(tmp_path / "other")[1].read_text(encoding="ascii")
    assert not audit.verify_run(tmp_path / "out", other)["ok"]
    m = tmp_path / "out" / "run_manifest.json"
    m.write_text(m.read_text(encoding="utf-8").replace('"operator": "', '"operator": "x'), encoding="utf-8")
    assert not audit.verify_run(tmp_path / "out", key)["ok"]


# ------------------------------------------------------------- tokens, profiles, rehydrate
def _tokens(key=None, profile="default", texts=("Priya Raman", "priya.raman@example.com", "4111 1111 1111 1111", "12/04/1985",
                                                 "2341 2341 2346")):
    cats = ["PERSON", "EMAIL_ADDRESS", "CREDIT_CARD", "DATE_OF_BIRTH", "IN_AADHAAR"]
    span = Span(id="s", file="f", text=" | ".join(texts), kind="paragraph")
    doc = Document(file="f", path="f", file_type="text", spans=[span])
    fs, pos = [], 0
    for t, c in zip(texts, cats):
        fs.append(Finding("s", "f", pos, pos + len(t), t, c, 0.9, "r", "L1 rules"))
        pos += len(t) + 3
    v = TokenVault(key=key, profile=profile)
    v.assign_all({"f": doc}, {"f": fs})
    return {f.entity_type: f.token for f in fs}, v


def test_keyed_tokens_are_the_same_in_every_run_and_need_the_key():
    a, _ = _tokens(key="k1")
    b, _ = _tokens(key="k1", texts=("Priya Raman", "priya.raman@example.com", "5500 0000 0000 0004", "01/01/1990", "2341 2341 2346"))
    c, _ = _tokens(key="k2")
    assert a["PERSON"] == b["PERSON"] and a["EMAIL_ADDRESS"] == b["EMAIL_ADDRESS"] and a["IN_AADHAAR"] == b["IN_AADHAAR"]
    assert a["CREDIT_CARD"] != b["CREDIT_CARD"] and a["PERSON"] != c["PERSON"]
    person_id = a["PERSON"][len("[PERSON_"):-1]
    assert len(person_id) == 8 and person_id in a["EMAIL_ADDRESS"]  # the e-mail still carries its owner's id
    plain, _ = _tokens()
    assert plain["PERSON"] == "[PERSON_001]" and plain["EMAIL_ADDRESS"] == "[EMAIL_001]"


def test_profiles_change_what_replaces_a_category():
    pci, _ = _tokens(profile="pci")
    assert pci["CREDIT_CARD"] == "[CARD_****1111]" and pci["PERSON"] == "[PERSON_001]"
    dpdp, _ = _tokens(profile="dpdp")
    assert dpdp["IN_AADHAAR"] == "[AADHAAR_****2346]" and dpdp["DATE_OF_BIRTH"] == "[DOB_1985]"


def test_rehydrate_puts_values_back_and_says_what_it_could_not():
    _, v = _tokens(profile="pci")
    v.values["[PERSON_001]"].add("Raman")
    v.values["[CARD_****1111]"].add("4111 2222 3333 1111")  # two cards behind one masked token
    text, restored, unknown = rehydrate("[PERSON_001] ([EMAIL_001]) paid with [CARD_****1111]; ask [PERSON_999]. [IMAGE WITHHELD]",
                                        v.values)
    assert text == "Priya Raman (priya.raman@example.com) paid with [CARD_****1111]; ask [PERSON_999]. [IMAGE WITHHELD]"
    assert restored == ["[PERSON_001]", "[EMAIL_001]"] and unknown == ["[CARD_****1111]", "[PERSON_999]"]


# ----------------------------------------------------------------------- new categories
def test_new_identifier_rules():
    from test_detect_units import best

    assert best("Beneficiary account 123456789012, IFSC HDFC0001234", "BANK_ACCOUNT")[2] >= 0.6
    assert best("pay priya.raman@okhdfcbank today", "UPI_ID")[1] == "priya.raman@okhdfcbank"
    assert best("National Insurance no. AB 12 34 56 C", "UK_NINO")[2] >= 0.6
    assert best("Voter ID ABC1234567 on file", "IN_VOTER_ID")[2] >= 0.35
    h = best("ticket ABC1234567 closed", "IN_VOTER_ID")
    assert h is None or h[2] < 0.35
    assert best("Driving licence no: MH12 20110012345", "DRIVING_LICENCE")[2] >= 0.6
    assert best("She was diagnosed with type 2 diabetes in 2019.", "HEALTH_DATA")[1] == "type 2 diabetes"
    assert best("Blood group: O+", "HEALTH_DATA")[2] >= 0.35
    h = best("We expect a positive outcome, grade A- or better.", "HEALTH_DATA")
    assert h is None or h[2] < 0.35
    h = best("Resides at Flat 4B, Sunrise Apartments, Sector 21, Gurugram 122001", "ADDRESS")
    assert h and h[1].startswith("Flat 4B") and h[1].endswith("122001")
    h = best("Section 16.5 applies to block 4 of the 2024 100200 batch", "ADDRESS")
    assert h is None or h[2] < 0.6 or "Section" not in h[1]
