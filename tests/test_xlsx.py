"""Spreadsheets: cells under headers, numbers stored as numbers, formulas, comments, hidden sheets,
sheet names and properties. Every planted value must be gone from the masked workbook (all parts)
and from the LLM text."""
import datetime as dt

import pytest

from test_leaks import surviving

PLANTED = ["Hans Müller", "kofi.mensah@example.com", "9820055512", "Priya Raman", "Siobhán Walsh", "Łukasz Nowak",
           "1985-04-12", "Marta Kowalczyk"]


def make_xlsx(path):
    import openpyxl
    from openpyxl.comments import Comment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(["Name", "Phone", "E-mail", "Date of birth", "Notes"])
    ws.append(["Hans Müller", 9820055512, "kofi.mensah@example.com", dt.date(1985, 4, 12), "vendor contact"])
    ws.append(["Łukasz Nowak", "+48 22 555 01 87", "lnowak@example.com", None, '=CONCAT("owner: ", "Marta Kowalczyk")'])
    ws["E2"].comment = Comment("Check with Siobhán Walsh before renewal", "Priya Raman")
    hidden = wb.create_sheet("Priya Raman")
    hidden.sheet_state = "hidden"
    hidden.append(["Reviewer", "Status"])
    hidden.append(["Siobhán Walsh", "open"])
    wb.properties.creator = "Priya Raman"
    wb.save(path)


@pytest.fixture(scope="module")
def xlsx_run(tmp_path_factory, settings):
    from optiv_pii_shield import run

    src = tmp_path_factory.mktemp("xlsx")
    make_xlsx(src / "contacts.xlsx")
    res = run([src / "contacts.xlsx"], settings, src / "out")
    res.out_dir, res.src = src / "out", src
    return res


def test_planted_values_present_in_source(xlsx_run):
    # the date is stored as an Excel serial number, so it is only checked in the LLM text
    in_file = [v for v in PLANTED if v != "1985-04-12"]
    assert {v for _, v in surviving(xlsx_run.src / "contacts.xlsx", in_file)} == set(in_file)


def test_extraction_reads_everything(xlsx_run):
    doc = xlsx_run.docs["contacts.xlsx"]
    texts = {s.text for s in doc.spans}
    assert {"Hans Müller", "9820055512", "1985-04-12", "Priya Raman", "Siobhán Walsh"} <= texts
    assert any(s.kind == "comment" for s in doc.spans) and doc.structure["hidden_sheets"] == 1
    cell = next(s for s in doc.spans if s.text == "9820055512")
    assert cell.header == "Phone" and cell.kind == "table_cell"


def test_masked_xlsx_written_and_clean(xlsx_run):
    assert not xlsx_run.errors, xlsx_run.errors
    out = xlsx_run.out_dir / "contacts.masked.xlsx"
    assert out.exists()
    left = surviving(out, PLANTED)
    assert not left, left


def test_llm_text_clean(xlsx_run):
    md = xlsx_run.redacted["contacts.xlsx"]
    assert not [v for v in PLANTED if v.lower() in md.lower()], md
    assert "| Name | Phone |" in md  # table structure kept


def test_masked_xlsx_still_opens(xlsx_run):
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_run.out_dir / "contacts.masked.xlsx")
    assert wb.worksheets[0]["A1"].value == "Name" and wb.worksheets[0]["A2"].value.startswith("[PERSON_")
    assert wb.worksheets[1].sheet_state == "hidden" and "Raman" not in wb.worksheets[1].title
